"""Tests for Core Studio Excel measurement export."""
from __future__ import annotations

import numpy as np
from openpyxl import load_workbook

from src.application import AppController
from src.domain.entities.core_layers import CoreDivision, CoreLayer
from src.domain.store import Store


def test_calibrated_core_excel_export_includes_depth(tmp_path) -> None:
    store = Store()
    controller = AppController(store)
    core_id = controller.create_core_entity(
        name="Calibrated",
        base_data=np.array([[[10, 20, 30]], [[40, 50, 60]]], dtype=np.uint8),
        mm_per_px=0.5,
    )
    path = tmp_path / "calibrated.xlsx"

    controller.export_core_to_excel(core_id, path)

    workbook = load_workbook(path, data_only=True)
    data = workbook["Core data"]
    assert [cell.value for cell in data[1]] == [
        "image_row_px", "layer_title", "depth_mm", "R_mean", "G_mean", "B_mean", "L_star", "a_star", "b_star"
    ]
    assert data["A2"].value == 0
    assert data["C3"].value == 0.5
    assert [data["D2"].value, data["E2"].value, data["F2"].value] == [10, 20, 30]
    assert workbook.sheetnames == ["Core data", "Layers"]


def test_uncalibrated_core_excel_export_omits_depth(tmp_path) -> None:
    store = Store()
    controller = AppController(store)
    core_id = controller.create_core_entity(
        name="Uncalibrated", base_data=np.zeros((2, 1, 3), dtype=np.uint8)
    )
    path = tmp_path / "uncalibrated.xlsx"

    controller.export_core_to_excel(core_id, path)

    workbook = load_workbook(path, data_only=True)
    data = workbook["Core data"]
    assert "depth_mm" not in [cell.value for cell in data[1]]
    assert workbook.sheetnames == ["Core data", "Layers"]


def test_excel_export_maps_layer_titles_and_details(tmp_path) -> None:
    store = Store()
    controller = AppController(store)
    division = CoreDivision("division", 2)
    core_id = controller.create_core_entity(
        name="Layered",
        base_data=np.zeros((4, 1, 3), dtype=np.uint8),
        mm_per_px=0.5,
    )
    core = store.get(core_id)
    core.divisions = [division]
    core.layers = [
        CoreLayer("upper", None, division.id, "Clay", "Fine laminated"),
        CoreLayer("lower", division.id, None, "Sand", "Coarse grains"),
    ]
    path = tmp_path / "layers.xlsx"

    controller.export_core_to_excel(core_id, path)

    workbook = load_workbook(path, data_only=True)
    data = workbook["Core data"]
    layers = workbook["Layers"]
    assert [data.cell(row, 2).value for row in range(2, 6)] == ["Clay", "Clay", "Sand", "Sand"]
    assert [cell.value for cell in layers[1]] == [
        "layer_title", "note", "start_row_px", "end_row_px", "start_depth_mm", "end_depth_mm"
    ]
    assert [cell.value for cell in layers[2]] == ["Clay", "Fine laminated", 0, 2, 0, 1]
    assert [cell.value for cell in layers[3]] == ["Sand", "Coarse grains", 2, 4, 1, 2]